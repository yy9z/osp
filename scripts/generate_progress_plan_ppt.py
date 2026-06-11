from pathlib import Path
import re
import glob

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path('/Users/yy/code/OSP')
OUT = ROOT / '院审汇报_已完成工作与后期计划_截图重点版.pptx'

IMG = {
    'logo': ROOT / 'docs/ppt_assets/ustc-logo-fixed.png',
    'cover': ROOT / 'campus-frontend/public/images/IMG_0014.jpg',
    's_dashboard': ROOT / 'docs/ppt_assets/screenshots/01_dashboard.png',
    's_agent': ROOT / 'docs/ppt_assets/screenshots/02_agent.png',
    's_secondhand': ROOT / 'docs/ppt_assets/screenshots/03_secondhand.png',
    's_lostfound': ROOT / 'docs/ppt_assets/screenshots/04_lostfound.png',
    's_messages': ROOT / 'docs/ppt_assets/screenshots/05_messages.png',
    's_dorm_info': ROOT / 'docs/ppt_assets/screenshots/06_dorm_info.png',
    's_dorm_repair': ROOT / 'docs/ppt_assets/screenshots/07_dorm_repair.png',
    's_navigation': ROOT / 'docs/ppt_assets/screenshots/08_navigation.png',
    's_profile': ROOT / 'docs/ppt_assets/screenshots/09_profile.png',
    'metrics': ROOT / 'docs/ppt_assets/project_metrics.png',
    'status': ROOT / 'docs/ppt_assets/engineering_status.png',
}


def collect_metrics():
    api = 0
    for p in glob.glob(str(ROOT / 'src/main/java/com/caspar/**/*.java'), recursive=True):
        if '/controller/' not in p:
            continue
        txt = Path(p).read_text(encoding='utf-8', errors='ignore')
        api += len(re.findall(r'@(Get|Post|Put|Delete|Patch)Mapping\s*\(', txt))

    counts = {
        'api': api,
        'controllers': len(glob.glob(str(ROOT / 'src/main/java/com/caspar/controller/*.java'))),
        'agent_classes': len(glob.glob(str(ROOT / 'src/main/java/com/caspar/agent/**/*.java'), recursive=True)),
        'views': len(glob.glob(str(ROOT / 'campus-frontend/src/views/**/*.vue'), recursive=True)),
        'mappers': len(glob.glob(str(ROOT / 'src/main/java/com/caspar/mapper/*.java'))),
    }
    return counts


M = collect_metrics()

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

DEEP = RGBColor(15, 23, 42)
TEXT = RGBColor(51, 65, 85)
BLUE = RGBColor(30, 64, 175)
WHITE = RGBColor(255, 255, 255)


def add_title(slide, text, top=0.38, size=34, color=DEEP):
    tb = slide.shapes.add_textbox(Inches(0.72), Inches(top), Inches(11.9), Inches(0.82))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color


def add_sub(slide, text, top=1.05, size=16, color=TEXT):
    tb = slide.shapes.add_textbox(Inches(0.72), Inches(top), Inches(11.9), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color


def add_cover_overlay(slide):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(8, 25, 48)
    rect.fill.transparency = 0.42
    rect.line.fill.background()


def add_pic(slide, key, x, y, w=None, h=None):
    path = IMG[key]
    if not path.exists():
        return
    kwargs = {}
    if w is not None:
        kwargs['width'] = Inches(w)
    if h is not None:
        kwargs['height'] = Inches(h)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def rounded(slide, x, y, w, h, fill=RGBColor(248, 250, 252), line=RGBColor(203, 213, 225)):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    return shp


# 1 Cover
s = prs.slides.add_slide(blank)
add_pic(s, 'cover', 0, 0, 13.333, 7.5)
add_cover_overlay(s)
if IMG['logo'].exists():
    add_pic(s, 'logo', 0.75, 0.45, h=0.85)

add_title(s, '院审汇报：已完成工作与后期计划', top=1.62, size=44, color=WHITE)
add_sub(s, '高校校园一站式平台（结合系统真实截图）', top=2.52, size=23, color=RGBColor(226, 232, 240))

info = s.shapes.add_textbox(Inches(0.82), Inches(4.95), Inches(6.2), Inches(1.9)).text_frame
for i, t in enumerate(['姓名：', '学号：', '指导老师：', '日期：2026年3月']):
    p = info.paragraphs[0] if i == 0 else info.add_paragraph()
    p.text = t
    p.font.size = Pt(23 if i == 0 else 20)
    p.font.color.rgb = WHITE

quote = s.shapes.add_textbox(Inches(7.0), Inches(5.35), Inches(5.8), Inches(1.2)).text_frame
quote.text = '本次重点：我们已经做成了什么，下一步将如何落地优化。'
quote.paragraphs[0].font.size = Pt(20)
quote.paragraphs[0].font.bold = True
quote.paragraphs[0].font.color.rgb = RGBColor(191, 219, 254)

# 2 Focus
s = prs.slides.add_slide(blank)
add_title(s, '1. 本次汇报聚焦点')
add_sub(s, '围绕“已完成工作”和“后期计划”两个核心问题展开')

left = rounded(s, 0.75, 1.6, 6.1, 5.35, fill=RGBColor(239, 246, 255), line=RGBColor(147, 197, 253))
tf = left.text_frame
tf.clear()
for i, t in enumerate([
    'A. 已完成工作（结合系统截图）',
    '• 智能助手：意图识别 + 槽位追问 + 卡片化结果',
    '• 二手交易：发布/浏览/收藏/消息联动闭环',
    '• 宿舍管理：宿舍信息 + 报修流程可运行',
    '• 校园导航：语义导航与地图路线展示',
    '• 基础能力：登录认证、权限控制、个人中心',
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = t
    p.font.size = Pt(24 if i == 0 else 16)
    p.font.bold = i == 0
    p.font.color.rgb = DEEP if i == 0 else TEXT

right = rounded(s, 7.1, 1.6, 5.5, 5.35, fill=RGBColor(236, 253, 245), line=RGBColor(134, 239, 172))
tf = right.text_frame
tf.clear()
for i, t in enumerate([
    'B. 后期计划（可执行）',
    '• 2026-03-16 至 2026-03-19：演示链路稳定化',
    '• 2026-03-20 至 2026-03-22：回归测试与细节打磨',
    '• 院审后 2-4 周：功能补全与性能优化',
    '',
    '目标：保证院审“讲得清、跑得通、问得住”。'
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = t
    p.font.size = Pt(24 if i == 0 else 16)
    p.font.bold = i == 0 or i == 5
    p.font.color.rgb = DEEP if i == 0 else TEXT

# 3 Completed overview
s = prs.slides.add_slide(blank)
add_title(s, '2. 已完成工作总览')
add_sub(s, '当前已完成可演示的核心模块（非PPT描述，均有系统页面支撑）')

rows = [
    ('智能助手（Agent）', '已完成核心链路', '自然语言输入 -> 意图识别 -> 工具调用 -> 卡片返回'),
    ('二手交易', '已形成闭环', '发布、列表、详情、收藏、消息沟通可联动'),
    ('宿舍管理', '已可稳定演示', '宿舍信息查看、报修申请与处理流程可用'),
    ('校园导航', '已具备语义能力', '地点语义识别、路线与时间展示、地图承接'),
    ('用户与安全', '已落地', 'JWT登录鉴权 + 角色权限控制 + 个人中心'),
]

col_x = [0.75, 3.9, 6.4]
col_w = [3.0, 2.3, 6.2]
for i, (c1, c2, c3) in enumerate([('模块', '完成状态', '说明')] + rows):
    y = 1.65 + i * 0.94
    is_head = i == 0
    bg = RGBColor(30, 64, 175) if is_head else RGBColor(248, 250, 252)
    fc = WHITE if is_head else DEEP
    for x, w, txt in zip(col_x, col_w, [c1, c2, c3]):
        c = rounded(s, x, y, w, 0.8, fill=bg, line=RGBColor(203, 213, 225))
        p = c.text_frame.paragraphs[0]
        p.text = txt
        p.font.size = Pt(15 if is_head else 14)
        p.font.bold = is_head or x == col_x[0]
        p.font.color.rgb = fc if is_head else (DEEP if x == col_x[0] else TEXT)
        p.alignment = PP_ALIGN.CENTER

# 4 Agent + Navigation
s = prs.slides.add_slide(blank)
add_title(s, '3. 已完成工作：智能助手与导航（真实截图）')
add_sub(s, '证据链：系统页面 + 可交互能力 + 可讲述场景')

frame1 = rounded(s, 0.75, 1.55, 6.1, 4.9)
frame2 = rounded(s, 7.0, 1.55, 5.6, 4.9)
add_pic(s, 's_agent', 0.88, 1.68, 5.85, 4.2)
add_pic(s, 's_navigation', 7.13, 1.68, 5.35, 4.2)

cap1 = s.shapes.add_textbox(Inches(0.9), Inches(5.95), Inches(5.8), Inches(0.45)).text_frame
cap1.text = '智能助手：支持多轮追问与结构化结果卡片'
cap1.paragraphs[0].font.size = Pt(14)
cap1.paragraphs[0].font.color.rgb = BLUE
cap1.paragraphs[0].font.bold = True

cap2 = s.shapes.add_textbox(Inches(7.2), Inches(5.95), Inches(5.2), Inches(0.45)).text_frame
cap2.text = '校园导航：语义地点识别 + 路线时长展示'
cap2.paragraphs[0].font.size = Pt(14)
cap2.paragraphs[0].font.color.rgb = BLUE
cap2.paragraphs[0].font.bold = True

# 5 Secondhand + Messages
s = prs.slides.add_slide(blank)
add_title(s, '4. 已完成工作：二手交易与消息联动（真实截图）')
add_sub(s, '证据链：交易流程页面 + 消息中心页面')

rounded(s, 0.75, 1.55, 8.25, 4.9)
rounded(s, 9.15, 1.55, 3.45, 4.9)
add_pic(s, 's_secondhand', 0.88, 1.68, 7.98, 4.2)
add_pic(s, 's_messages', 9.28, 1.68, 3.2, 4.2)

note = rounded(s, 0.75, 6.55, 11.85, 0.75, fill=RGBColor(239, 246, 255), line=RGBColor(147, 197, 253))
p = note.text_frame.paragraphs[0]
p.text = '已完成：发布/筛选/收藏/详情/会话沟通主链路，支持院审现场完整演示。'
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = RGBColor(30, 64, 175)
p.alignment = PP_ALIGN.CENTER

# 6 Dormitory + LostFound + Profile
s = prs.slides.add_slide(blank)
add_title(s, '5. 已完成工作：宿舍、失物招领与个人中心（真实截图）')
add_sub(s, '证据链：报修、招领、个人信息维护页面均可访问')

cards = [
    ('宿舍报修', 's_dorm_repair'),
    ('失物招领', 's_lostfound'),
    ('个人中心', 's_profile'),
]
for i, (name, key) in enumerate(cards):
    x = 0.75 + i * 4.2
    rounded(s, x, 1.55, 3.85, 5.75)
    add_pic(s, key, x + 0.12, 1.68, 3.6, 4.95)
    t = s.shapes.add_textbox(Inches(x + 0.12), Inches(6.73), Inches(3.6), Inches(0.35)).text_frame
    t.text = name
    t.paragraphs[0].font.size = Pt(15)
    t.paragraphs[0].font.bold = True
    t.paragraphs[0].font.color.rgb = BLUE
    t.paragraphs[0].alignment = PP_ALIGN.CENTER

# 7 Stage result
s = prs.slides.add_slide(blank)
add_title(s, '6. 阶段性成果（量化 + 工程验证）')
add_sub(s, '展示“做了什么”之外，也说明“工程上是否站得住”')

add_pic(s, 'metrics', 0.75, 1.55, 5.95, 3.25)
add_pic(s, 'status', 0.75, 4.95, 5.95, 2.2)

panel = rounded(s, 6.95, 1.55, 5.65, 5.6, fill=RGBColor(255, 255, 255), line=RGBColor(203, 213, 225))
tf = panel.text_frame
tf.clear()
lines = [
    '项目规模快照',
    f'• API 映射：{M["api"]}',
    f'• Agent 相关类：{M["agent_classes"]}',
    f'• 控制器：{M["controllers"]}',
    f'• 前端视图：{M["views"]}',
    f'• Mapper：{M["mappers"]}',
    '',
    '可验证状态',
    '• 后端编译通过（Maven）',
    '• 前端构建通过（Vite）',
    '',
    '结论：当前系统具备院审所需的可运行与可演示基础。'
]
for i, t in enumerate(lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = t
    p.font.size = Pt(23 if i in [0, 7] else 15)
    p.font.bold = i in [0, 7, 11]
    p.font.color.rgb = DEEP if i in [0, 7] else TEXT

# 8 Future plan timeline
s = prs.slides.add_slide(blank)
add_title(s, '7. 后期计划（按时间推进）')
add_sub(s, '围绕“院审前稳定演示 + 院审后持续优化”两条线并行')

phases = [
    ('阶段1', '2026-03-16 至 2026-03-19', ['冻结院审演示脚本', '重点修复主链路边界问题', '完善异常提示与兜底提示']),
    ('阶段2', '2026-03-20 至 2026-03-22', ['执行核心场景回归（报修/二手/导航）', '优化前端首屏与模块加载体验', '整理答辩问答素材与演示备份方案']),
    ('阶段3', '院审后 2-4 周', ['补强失物招领完整闭环', '完善系统化测试与日志观测', '继续优化跨模块协同与性能']),
]
for i, (stage, date, tasks) in enumerate(phases):
    y = 1.6 + i * 1.86
    tag = rounded(s, 0.75, y, 1.7, 1.45, fill=RGBColor(30, 64, 175), line=RGBColor(30, 64, 175))
    p = tag.text_frame.paragraphs[0]
    p.text = stage
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    box = rounded(s, 2.62, y, 9.95, 1.45, fill=RGBColor(248, 250, 252), line=RGBColor(203, 213, 225))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = date
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BLUE
    for task in tasks:
        p = tf.add_paragraph()
        p.text = '• ' + task
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT

# 9 Risks
s = prs.slides.add_slide(blank)
add_title(s, '8. 风险识别与应对方案')
add_sub(s, '提前暴露风险并给出动作，提升院审答辩说服力')

risk_rows = [
    ('潜在风险', '可能影响', '已规划应对'),
    ('演示链路中断或异常', '现场展示节奏被打断', '预设演示数据 + 备用路径 + 本地备份截图'),
    ('高峰场景性能波动', '页面响应不稳定', '优先优化高频页面与接口，减少首屏负担'),
    ('跨模块状态一致性', '评委追问细节时风险增大', '强化关键链路回归与日志检查'),
]
xx = [0.75, 4.45, 7.75]
ww = [3.5, 3.1, 4.85]
for i, row in enumerate(risk_rows):
    y = 1.8 + i * 1.25
    for j, txt in enumerate(row):
        is_head = i == 0
        cell = rounded(s, xx[j], y, ww[j], 1.0,
                       fill=RGBColor(30, 64, 175) if is_head else RGBColor(248, 250, 252),
                       line=RGBColor(203, 213, 225))
        p = cell.text_frame.paragraphs[0]
        p.text = txt
        p.font.size = Pt(15 if is_head else 14)
        p.font.bold = is_head or j == 0
        p.font.color.rgb = WHITE if is_head else (DEEP if j == 0 else TEXT)
        p.alignment = PP_ALIGN.CENTER

commit = rounded(s, 0.75, 6.35, 11.85, 0.8, fill=RGBColor(220, 252, 231), line=RGBColor(134, 239, 172))
p = commit.text_frame.paragraphs[0]
p.text = '承诺：院审前以“稳定演示 + 关键链路闭环”为最高优先级。'
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = RGBColor(20, 83, 45)
p.alignment = PP_ALIGN.CENTER

# 10 Ending
s = prs.slides.add_slide(blank)
add_title(s, '9. 总结与预期结果', top=1.2, size=40, color=DEEP)
add_sub(s, '通过“已完成工作 + 后期计划”双线呈现，确保院审可讲、可演、可追问。', top=2.0, size=19, color=TEXT)

sum_box = rounded(s, 0.9, 2.7, 11.6, 3.35, fill=RGBColor(239, 246, 255), line=RGBColor(147, 197, 253))
tf = sum_box.text_frame
tf.clear()
for i, t in enumerate([
    '本次可清晰回答三个问题：',
    '1) 已完成了什么：核心模块已有真实页面与可运行链路。',
    '2) 结果如何：已有可量化规模与可验证工程状态。',
    '3) 下一步怎么做：已按日期拆解出可执行优化计划。',
    '',
    '预期：在院审汇报中更稳地展示项目完成度与项目管理能力。'
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = t
    p.font.size = Pt(24 if i == 0 else 18)
    p.font.bold = i in [0, 5]
    p.font.color.rgb = DEEP if i == 0 else TEXT

thanks = s.shapes.add_textbox(Inches(0.95), Inches(6.45), Inches(6.0), Inches(0.45)).text_frame
thanks.text = '谢谢各位老师，请批评指正'
thanks.paragraphs[0].font.size = Pt(24)
thanks.paragraphs[0].font.bold = True
thanks.paragraphs[0].font.color.rgb = BLUE

qa = s.shapes.add_textbox(Inches(10.2), Inches(6.38), Inches(2.4), Inches(0.55)).text_frame
qa.text = 'Q & A'
qa.paragraphs[0].font.size = Pt(32)
qa.paragraphs[0].font.bold = True
qa.paragraphs[0].font.color.rgb = RGBColor(30, 64, 175)
qa.paragraphs[0].alignment = PP_ALIGN.RIGHT

prs.save(str(OUT))
print(OUT)
