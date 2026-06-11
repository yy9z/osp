from pathlib import Path
import re
import glob
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path('/Users/yy/code/OSP')
OUT = ROOT / '项目说明_阶段成果与后续计划_可直接使用版.pptx'

IMG = {
    'cover_bg': ROOT / 'campus-frontend/public/images/IMG_0014.jpg',
    'cover_bg_alt': ROOT / 'campus-frontend/public/images/campus-bg1.jpg',
    'login': ROOT / 'docs/ppt_assets/keyshots/01_login.png',
    'dashboard': ROOT / 'docs/ppt_assets/keyshots/02_dashboard.png',
    'agent': ROOT / 'docs/ppt_assets/screenshots/02_agent.png',
    'secondhand': ROOT / 'docs/ppt_assets/keyshots/04_secondhand.png',
    'lostfound': ROOT / 'docs/ppt_assets/keyshots/05_lostfound.png',
    'messages': ROOT / 'docs/ppt_assets/keyshots/06_messages.png',
}


def metrics():
    api = 0
    for p in glob.glob(str(ROOT / 'src/main/java/com/caspar/**/*.java'), recursive=True):
        if '/controller/' not in p:
            continue
        txt = Path(p).read_text(encoding='utf-8', errors='ignore')
        api += len(re.findall(r'@(Get|Post|Put|Delete|Patch)Mapping\s*\(', txt))
    return {
        'api': api,
        'controllers': len(glob.glob(str(ROOT / 'src/main/java/com/caspar/controller/*.java'))),
        'agent_classes': len(glob.glob(str(ROOT / 'src/main/java/com/caspar/agent/**/*.java'), recursive=True)),
        'views': len(glob.glob(str(ROOT / 'campus-frontend/src/views/**/*.vue'), recursive=True)),
    }

M = metrics()

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

C_BG = RGBColor(250, 252, 255)
C_TITLE = RGBColor(15, 23, 42)
C_TEXT = RGBColor(51, 65, 85)
C_BLUE = RGBColor(30, 64, 175)
C_WHITE = RGBColor(255, 255, 255)


def base_bg(slide, mood='blue'):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = C_BG
    r.line.fill.background()

    # Minimal header ribbon
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(1.45))
    top.fill.solid(); top.fill.fore_color.rgb = RGBColor(245, 248, 255)
    top.line.fill.background()

    # Accent line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.42), prs.slide_width, Inches(0.03))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(191, 219, 254)
    line.line.fill.background()


def title(slide, t, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12.0), Inches(0.9)).text_frame
    p = tb.paragraphs[0]
    p.text = t
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = C_TITLE
    if sub:
        sb = slide.shapes.add_textbox(Inches(0.72), Inches(1.03), Inches(11.9), Inches(0.45)).text_frame
        p = sb.paragraphs[0]
        p.text = sub
        p.font.size = Pt(16)
        p.font.color.rgb = C_TEXT


def card(slide, x, y, w, h, fill=RGBColor(255,255,255), line=RGBColor(203,213,225)):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    return s


def screenshot_slide(t, sub, img_key, bullets):
    s = prs.slides.add_slide(blank)
    base_bg(s, mood='blue')
    title(s, t, sub)
    # Keep screenshot in near-16:9 ratio to avoid squashing
    frame = card(s, 1.55, 1.5, 10.25, 5.25, fill=RGBColor(255, 255, 255), line=RGBColor(203, 213, 225))
    img = IMG[img_key]
    if img.exists():
        s.shapes.add_picture(str(img), Inches(2.18), Inches(1.6), width=Inches(8.98), height=Inches(5.05))
    note = card(s, 1.55, 6.82, 10.25, 0.52, fill=RGBColor(239,246,255), line=RGBColor(147,197,253))
    tf = note.text_frame
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ('• ' + b) if i > 0 else b
        p.font.size = Pt(11 if i > 0 else 12)
        p.font.bold = i == 0
        p.font.color.rgb = C_BLUE if i == 0 else C_TEXT

# 1 Cover
s = prs.slides.add_slide(blank)

# Cover background style v2: bright campus photo + left info panel
bg_path = IMG['cover_bg_alt'] if IMG['cover_bg_alt'].exists() else IMG['cover_bg']
if bg_path.exists():
    s.shapes.add_picture(str(bg_path), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
    overlay = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(255, 255, 255)
    overlay.fill.transparency = 0.18
    overlay.line.fill.background()
else:
    base_bg(s)

# Left information panel
panel = card(s, 0.75, 0.75, 6.0, 6.0, fill=RGBColor(19, 60, 128), line=RGBColor(19, 60, 128))
panel.fill.transparency = 0.09

line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.15), Inches(3.45), Inches(3.0), Inches(0.05))
line.fill.solid(); line.fill.fore_color.rgb = RGBColor(147, 197, 253)
line.line.fill.background()

h = s.shapes.add_textbox(Inches(1.15), Inches(1.9), Inches(5.35), Inches(1.9)).text_frame
p = h.paragraphs[0]
p.text = '一站式校园事务智能协同平台'
p.font.size = Pt(44); p.font.bold = True; p.font.color.rgb = C_WHITE
p = h.add_paragraph()
p.text = '阶段成果：已完成工作与下一步计划'
p.font.size = Pt(22); p.font.bold = False; p.font.color.rgb = RGBColor(219, 234, 254)

chip = card(s, 1.15, 1.15, 2.2, 0.42, fill=RGBColor(30, 64, 175), line=RGBColor(147, 197, 253))
chip.fill.transparency = 0.0
cp = chip.text_frame.paragraphs[0]
cp.text = 'PROJECT REPORT'
cp.font.size = Pt(12)
cp.font.bold = True
cp.font.color.rgb = RGBColor(219, 234, 254)
cp.alignment = PP_ALIGN.CENTER

info = s.shapes.add_textbox(Inches(1.15), Inches(4.15), Inches(5.35), Inches(1.8)).text_frame
for i, t in enumerate(['姓名：', '学号：', '日期：2026 年 3 月']):
    p = info.paragraphs[0] if i == 0 else info.add_paragraph()
    p.text = t
    p.font.size = Pt(19 if i == 0 else 17)
    p.font.color.rgb = C_WHITE

# 2 Agenda
s = prs.slides.add_slide(blank)
base_bg(s, mood='green')
title(s, '1. 内容概览', '围绕阶段成果、系统架构与后续执行计划展开')

left = card(s, 0.8, 1.9, 5.7, 4.9, fill=RGBColor(239,246,255), line=RGBColor(147,197,253))
right = card(s, 6.85, 1.9, 5.7, 4.9, fill=RGBColor(236,253,245), line=RGBColor(134,239,172))

for box, lines in [
    (left, ['A. 阶段成果概览', '核心模块完成情况', '关键页面真实截图', '可稳定演示的主链路']),
    (right, ['B. 架构与执行计划', '系统运行逻辑', '后续工作安排', '验收标准与推进方式'])
]:
    tf = box.text_frame; tf.clear()
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = t if i == 0 else '• ' + t
        p.font.size = Pt(24 if i == 0 else 16)
        p.font.bold = i == 0
        p.font.color.rgb = C_TITLE if i == 0 else C_TEXT

# 3 Completed summary
s = prs.slides.add_slide(blank)
base_bg(s, mood='blue')
title(s, '2. 阶段成果概览', '核心业务模块已上线，具备稳定演示基础')

rows = [
    ('模块', '完成状态', '说明'),
    ('登录与权限', '已完成', '统一身份认证 + JWT 鉴权 + 角色权限控制'),
    ('首页聚合', '已完成', '模块入口、动态信息、任务提醒等聚合展示'),
    ('智能助手', '已完成', '自然语言入口，支持多场景引导与任务承接'),
    ('二手交易', '已完成', '发布、检索、收藏、沟通链路可运行'),
    ('失物招领', '基础能力已完成', '公告发布、列表检索与个人视图已可用'),
    ('消息中心', '已完成', '会话列表、历史消息与未读提醒已可用'),
]
xx=[0.72,3.8,6.1]; ww=[2.9,2.2,6.62]
for i,row in enumerate(rows):
    y=1.65+i*0.8
    for j,txt in enumerate(row):
        head=i==0
        c=card(s,xx[j],y,ww[j],0.68,fill=RGBColor(30,64,175) if head else RGBColor(255,255,255))
        p=c.text_frame.paragraphs[0]
        p.text=txt
        p.font.size=Pt(14 if head else 13)
        p.font.bold=head or j==0
        p.font.color.rgb=C_WHITE if head else (C_TITLE if j==0 else C_TEXT)
        p.alignment=PP_ALIGN.CENTER

# 4 Architecture (simplified)
s = prs.slides.add_slide(blank)
base_bg(s, mood='green')
title(s, '3. 系统运行逻辑（简化架构）', '用户请求经 Agent 编排后，由业务模块执行并返回结构化结果')

steps = [
    ('1. 交互层（前端）', ['接收用户请求并展示结果', '入口覆盖报修、交易、导航等场景'], RGBColor(239,246,255)),
    ('2. 编排层（Agent）', ['识别意图与关键参数', '将自然语言转为可执行任务'], RGBColor(236,253,245)),
    ('3. 业务层（模块服务）', ['调用宿舍、二手、失物、消息、导航模块', '按业务规则执行并返回结果'], RGBColor(255,247,237)),
    ('4. 数据层（存储与服务）', ['MySQL/Redis提供数据与会话支撑', '外部服务补充地图与文件能力'], RGBColor(243,232,255)),
]

x0 = 0.75
w = 2.9
gap = 0.23
y = 2.05
h = 3.6
for i, (name, lines, fill) in enumerate(steps):
    x = x0 + i * (w + gap)
    b = card(s, x, y, w, h, fill=fill)
    tf = b.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = name
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_TITLE
    for line in lines:
        p = tf.add_paragraph()
        p.text = '• ' + line
        p.font.size = Pt(14)
        p.font.color.rgb = C_TEXT

    if i < len(steps) - 1:
        arr = s.shapes.add_shape(
            MSO_SHAPE.CHEVRON,
            Inches(x + w + 0.03),
            Inches(y + 1.45),
            Inches(0.18),
            Inches(0.7),
        )
        arr.fill.solid()
        arr.fill.fore_color.rgb = RGBColor(147, 197, 253)
        arr.line.fill.background()

tip = card(s, 0.75, 6.05, 12.0, 1.2, fill=RGBColor(219,234,254), line=RGBColor(147,197,253))
tf = tip.text_frame
tf.clear()
p = tf.paragraphs[0]
p.text = '架构主线：前端发起请求，Agent负责任务编排，业务模块执行，数据层提供支撑。'
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = RGBColor(30,64,175)
p.alignment = PP_ALIGN.CENTER
p = tf.add_paragraph()
p.text = '核心价值：在保留原系统稳定性的前提下，实现跨模块统一入口与协同处理。'
p.font.size = Pt(14)
p.font.color.rgb = C_TEXT
p.alignment = PP_ALIGN.CENTER

# Screenshot slides (one screenshot each)
screenshot_slide('4. 阶段成果：登录与权限体系', '统一身份认证与权限控制已落地', 'login', [
    '重点：登录入口稳定，账号认证流程完整',
    '按角色控制菜单展示与接口访问范围',
])

screenshot_slide('5. 阶段成果：首页聚合能力', '多模块信息在首页统一呈现，支持快速触达', 'dashboard', [
    '重点：首页承接交易、失物、消息、宿舍、导航等入口',
    '可作为演示起点，直观体现平台联动能力',
])

screenshot_slide('6. 阶段成果：智能助手入口', '自然语言入口已接入核心业务流程', 'agent', [
    '重点：可通过一句话发起校园事务请求',
    '已支持报修、交易、失物、导航等场景引导',
])

screenshot_slide('7. 阶段成果：二手交易模块', '已形成从浏览到沟通的业务闭环', 'secondhand', [
    '重点：商品检索、分类筛选、推荐展示已可用',
    '与消息中心联动，支持后续沟通与成交推进',
])

screenshot_slide('8. 阶段成果：失物招领模块', '支持信息发布、筛选检索与状态流转', 'lostfound', [
    '重点：寻物/招领双视图与个人记录页面已落地',
    '已具备继续完善完整闭环的基础',
])

screenshot_slide('9. 阶段成果：消息中心模块', '系统内沟通链路已打通', 'messages', [
    '重点：会话列表、聊天记录、发送链路已可运行',
    '可承接二手交易等跨模块沟通场景',
])

# 10 Plan
s = prs.slides.add_slide(blank)
base_bg(s, mood='green')
title(s, '10. 后续执行计划', '推进主线：演示稳定化 -> 功能闭环补齐 -> 质量与性能优化')

phases=[
 ('阶段一', '演示链路稳定化', ['固化 3 条主演示路径（登录、助手、交易/消息）', '排查阻断问题并统一异常提示', '验收：主链路可连续完整演示']),
 ('阶段二', '功能闭环补齐', ['完善失物招领关键流程（发布/查看/状态）', '补强报修与消息联动反馈一致性', '验收：关键流程端到端可走通']),
 ('阶段三', '质量与性能优化', ['建立核心流程回归检查清单', '优化高频页面加载与交互响应', '验收：演示过程稳定、体验流畅'])
]
for i,(a,b,cs) in enumerate(phases):
    y=1.55+i*1.9
    tag=card(s,0.75,y,2.4,1.45,fill=RGBColor(30,64,175),line=RGBColor(30,64,175))
    p=tag.text_frame.paragraphs[0]; p.text=a; p.font.size=Pt(16); p.font.bold=True; p.font.color.rgb=C_WHITE; p.alignment=PP_ALIGN.CENTER
    box=card(s,3.3,y,9.4,1.45)
    tf=box.text_frame; tf.clear()
    p=tf.paragraphs[0]; p.text=b; p.font.size=Pt(16); p.font.bold=True; p.font.color.rgb=C_BLUE
    for c in cs:
        p=tf.add_paragraph(); p.text='• '+c; p.font.size=Pt(13); p.font.color.rgb=C_TEXT

# 11 ending
s = prs.slides.add_slide(blank)
base_bg(s)

# subtle decorative shapes for a cleaner visual ending
blob1 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-0.8), Inches(4.8), Inches(3.2), Inches(3.2))
blob1.fill.solid(); blob1.fill.fore_color.rgb = RGBColor(219, 234, 254); blob1.fill.transparency = 0.45
blob1.line.fill.background()
blob2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.0), Inches(-0.8), Inches(3.4), Inches(3.4))
blob2.fill.solid(); blob2.fill.fore_color.rgb = RGBColor(224, 242, 254); blob2.fill.transparency = 0.45
blob2.line.fill.background()

title_box = s.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.8), Inches(0.9)).text_frame
p = title_box.paragraphs[0]
p.text = '总结与后续重点'
p.font.size = Pt(42)
p.font.bold = True
p.font.color.rgb = C_TITLE
p.alignment = PP_ALIGN.CENTER

sub = s.shapes.add_textbox(Inches(0.8), Inches(1.95), Inches(11.8), Inches(0.5)).text_frame
p = sub.paragraphs[0]
p.text = '当前版本已具备稳定可用能力，下一步聚焦闭环完善与体验提升。'
p.font.size = Pt(17)
p.font.color.rgb = C_TEXT
p.alignment = PP_ALIGN.CENTER

panel = card(s, 1.05, 2.55, 11.2, 3.2, fill=RGBColor(255, 255, 255), line=RGBColor(191, 219, 254))

rows = [
    ('成果现状', '核心业务模块已可用，关键页面可直接展示。', RGBColor(30, 64, 175), RGBColor(239, 246, 255)),
    ('当前优势', '统一入口与模块协同主线清晰，系统逻辑易于理解。', RGBColor(22, 101, 52), RGBColor(236, 253, 245)),
    ('后续方向', '优先补齐关键闭环细节，并持续优化稳定性与响应体验。', RGBColor(154, 52, 18), RGBColor(255, 247, 237)),
]

for i, (k, v, kcolor, fill) in enumerate(rows):
    y = 2.82 + i * 0.98
    tag = card(s, 1.45, y, 1.75, 0.62, fill=fill, line=RGBColor(203, 213, 225))
    tf = tag.text_frame
    tf.clear()
    pp = tf.paragraphs[0]
    pp.text = k
    pp.font.size = Pt(16)
    pp.font.bold = True
    pp.font.color.rgb = kcolor
    pp.alignment = PP_ALIGN.CENTER

    body = s.shapes.add_textbox(Inches(3.45), Inches(y + 0.08), Inches(8.5), Inches(0.52)).text_frame
    bb = body.paragraphs[0]
    bb.text = v
    bb.font.size = Pt(16)
    bb.font.color.rgb = C_TEXT

line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.9), Inches(6.18), Inches(9.55), Inches(0.03))
line.fill.solid()
line.fill.fore_color.rgb = RGBColor(191, 219, 254)
line.line.fill.background()

end = s.shapes.add_textbox(Inches(0.8), Inches(6.28), Inches(11.8), Inches(0.45)).text_frame
pp = end.paragraphs[0]
pp.text = '感谢各位老师和同学的指导'
pp.font.size = Pt(24)
pp.font.bold = True
pp.font.color.rgb = C_BLUE
pp.alignment = PP_ALIGN.CENTER

prs.save(str(OUT))
print(OUT)
