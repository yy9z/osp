from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pathlib import Path

ROOT = Path('/Users/yy/code/OSP')
OUT = ROOT / '院审汇报_高校校园一站式平台_图文增强版.pptx'

IMG = {
    'campus_cover': ROOT / 'campus-frontend/public/images/IMG_0014.jpg',
    'campus_bg': ROOT / 'campus-frontend/public/images/campus-bg1.jpg',
    'building': ROOT / 'campus-frontend/public/images/building-bg.jpg',
    'library': ROOT / 'campus-frontend/public/images/library-bg.jpg',
    'map': ROOT / 'docs/ppt_assets/campus-map-fixed.png',
    'laptop': ROOT / 'campus-frontend/public/images/laptop.jpg',
    'phone': ROOT / 'campus-frontend/public/images/phone.jpg',
    'book': ROOT / 'campus-frontend/public/images/book.jpg',
    'logo': ROOT / 'docs/ppt_assets/ustc-logo-fixed.png',
    'arch': ROOT / 'docs/ppt_assets/architecture_overview.png',
    'flow': ROOT / 'docs/ppt_assets/agent_workflow.png',
    'metrics': ROOT / 'docs/ppt_assets/project_metrics.png',
    'status': ROOT / 'docs/ppt_assets/engineering_status.png',
    'shot_dashboard': ROOT / 'docs/ppt_assets/screenshots/01_dashboard.png',
    'shot_agent': ROOT / 'docs/ppt_assets/screenshots/02_agent.png',
    'shot_secondhand': ROOT / 'docs/ppt_assets/screenshots/03_secondhand.png',
    'shot_lostfound': ROOT / 'docs/ppt_assets/screenshots/04_lostfound.png',
    'shot_messages': ROOT / 'docs/ppt_assets/screenshots/05_messages.png',
    'shot_dorm_info': ROOT / 'docs/ppt_assets/screenshots/06_dorm_info.png',
    'shot_dorm_repair': ROOT / 'docs/ppt_assets/screenshots/07_dorm_repair.png',
    'shot_navigation': ROOT / 'docs/ppt_assets/screenshots/08_navigation.png',
    'shot_profile': ROOT / 'docs/ppt_assets/screenshots/09_profile.png',
}

BLUE = RGBColor(23, 79, 156)
DEEP = RGBColor(15, 23, 42)
TEXT = RGBColor(51, 65, 85)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def add_bg(slide, path):
    slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)

def add_dark_overlay(slide, alpha=0.35):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(8, 25, 48)
    rect.fill.transparency = alpha
    rect.line.fill.background()

def add_title(slide, text, left=0.7, top=0.35, width=12.0, size=36, color=DEEP):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    return tb

def add_subtitle(slide, text, left=0.72, top=1.15, width=11.8, size=17, color=TEXT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return tb

def add_card(slide, x, y, w, h, title, lines, fill_rgb):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = fill_rgb
    card.line.color.rgb = RGBColor(203, 213, 225)

    tf = card.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = DEEP

    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT

# Slide 1: Cover
slide = prs.slides.add_slide(blank)
add_bg(slide, IMG['campus_cover'])
add_dark_overlay(slide, 0.42)
if IMG['logo'].exists():
    slide.shapes.add_picture(str(IMG['logo']), Inches(0.75), Inches(0.45), height=Inches(0.9))

add_title(slide, '基于 Agent 的高校校园智能事务处理平台', left=0.75, top=1.55, width=11.8, size=42, color=WHITE)
add_subtitle(slide, '院审汇报（图文增强版）', left=0.78, top=2.45, width=7.0, size=23, color=RGBColor(226, 232, 240))

info = slide.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(6.5), Inches(2.1)).text_frame
for i, txt in enumerate(['姓名：', '学号：', '指导老师：', '日期：2026 年 3 月']):
    p = info.paragraphs[0] if i == 0 else info.add_paragraph()
    p.text = txt
    p.font.size = Pt(24 if i == 0 else 21)
    p.font.color.rgb = WHITE

quote = slide.shapes.add_textbox(Inches(7.0), Inches(5.3), Inches(5.6), Inches(1.5)).text_frame
p = quote.paragraphs[0]
p.text = '一句话定位：用自然语言统一校园高频事务入口'
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = RGBColor(191, 219, 254)

# Slide 2: Review focus
slide = prs.slides.add_slide(blank)
add_title(slide, '1. 项目定位与院审关注点', size=34)
add_subtitle(slide, '从“功能堆叠”升级为“统一智能入口”，强调可用性、可扩展性与可演示性', size=16)

if IMG['campus_bg'].exists():
    slide.shapes.add_picture(str(IMG['campus_bg']), Inches(7.2), Inches(1.4), Inches(5.8), Inches(5.7))
    ov = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.2), Inches(1.4), Inches(5.8), Inches(5.7))
    ov.fill.solid()
    ov.fill.fore_color.rgb = RGBColor(15, 23, 42)
    ov.fill.transparency = 0.5
    ov.line.fill.background()

cards = [
    ('痛点', ['入口分散，用户需要在多个页面间跳转', '传统表单流程复杂，沟通成本高'], RGBColor(239, 246, 255)),
    ('方案', ['Agent 统一入口：意图识别 + 槽位补全 + 工具编排', '返回结构化卡片，支持连续追问'], RGBColor(236, 253, 245)),
    ('价值', ['降低新用户学习成本', '提升事务办理效率与完成率'], RGBColor(255, 247, 237)),
]
for i, (t, l, c) in enumerate(cards):
    add_card(slide, 0.75, 1.65 + i * 1.72, 6.1, 1.5, t, l, c)

# Slide 3: Goals
slide = prs.slides.add_slide(blank)
add_title(slide, '2. 建设目标：可讲、可演、可扩展', size=34)

left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.4), Inches(6.2), Inches(5.7))
left.fill.solid()
left.fill.fore_color.rgb = RGBColor(248, 250, 252)
left.line.color.rgb = RGBColor(203, 213, 225)
tf = left.text_frame
tf.clear()
for idx, txt in enumerate([
    '核心目标',
    '1) 自然语言即可发起校园事务',
    '2) 多轮追问补全参数，减少失败请求',
    '3) 业务结果结构化展示，提升可解释性',
    '4) 保持原有业务系统，避免推倒重做',
    '',
    '评审可见成果',
    '• 覆盖宿舍、二手、失物、导航、消息等高频场景',
    '• 支持端到端演示链路（输入 -> 处理 -> 卡片结果）',
    '• 项目具备后续横向扩展能力（新增工具即新增能力）'
]):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.text = txt
    p.font.size = Pt(24 if idx == 0 else 17)
    p.font.bold = (idx == 0 or idx == 6)
    p.font.color.rgb = DEEP if (idx == 0 or idx == 6) else TEXT

if IMG['map'].exists():
    slide.shapes.add_picture(str(IMG['map']), Inches(7.2), Inches(1.4), Inches(5.5), Inches(2.65))
if IMG['building'].exists():
    slide.shapes.add_picture(str(IMG['building']), Inches(7.2), Inches(4.15), Inches(2.65), Inches(2.9))
if IMG['laptop'].exists():
    slide.shapes.add_picture(str(IMG['laptop']), Inches(10.05), Inches(4.15), Inches(2.65), Inches(2.9))

# Slide 4: Architecture
slide = prs.slides.add_slide(blank)
add_title(slide, '3. 总体架构：前后端分离 + Agent 编排中枢', size=34)
add_subtitle(slide, '前端负责交互承接，后端负责业务能力，Agent 负责跨模块理解与调度', size=16)
if IMG['arch'].exists():
    slide.shapes.add_picture(str(IMG['arch']), Inches(0.6), Inches(1.25), Inches(12.2), Inches(5.35))

chips = [
    ('技术底座：Spring Boot + MyBatis + MySQL + Vue3 + Element Plus', 0.8),
    ('安全机制：JWT + Spring Security + 角色权限控制', 4.6),
    ('会话管理：Redis + 多轮上下文', 8.8),
]
for txt, x in chips:
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(6.75), Inches(3.8), Inches(0.45))
    chip.fill.solid(); chip.fill.fore_color.rgb = RGBColor(219, 234, 254)
    chip.line.color.rgb = RGBColor(147, 197, 253)
    p = chip.text_frame.paragraphs[0]
    p.text = txt
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(30, 64, 175)
    p.alignment = PP_ALIGN.CENTER

# Slide 5: Agent workflow
slide = prs.slides.add_slide(blank)
add_title(slide, '4. Agent 创新流程：从“问答”到“可执行任务”', size=33)
add_subtitle(slide, '关键能力：意图识别、槽位补全、工具编排、结果卡片、后续建议', size=16)
if IMG['flow'].exists():
    slide.shapes.add_picture(str(IMG['flow']), Inches(0.55), Inches(1.4), Inches(12.2), Inches(3.45))

for i, (title, body) in enumerate([
    ('多轮追问', '缺失槽位时自动追问，减少用户二次填写成本'),
    ('跨模块调用', '统一入口可调度报修、二手、失物、导航、消息工具'),
    ('结构化输出', '通过卡片承接结果，支持一键跳转下一步操作'),
]):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*4.2), Inches(5.1), Inches(3.9), Inches(1.95))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(248, 250, 252)
    box.line.color.rgb = RGBColor(203, 213, 225)
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = DEEP
    p = tf.add_paragraph()
    p.text = body
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT

# Slide 6: Core scenarios
slide = prs.slides.add_slide(blank)
add_title(slide, '5. 三个高频场景：能落地、能演示、能答辩', size=34)

scenario = [
    ('宿舍事务（真实截图）', IMG['shot_dorm_repair'], ['自然语言报修', '工单状态流转', '宿管处理闭环']),
    ('二手交易（真实截图）', IMG['shot_secondhand'], ['发布/筛选/收藏', '消息联动沟通', '交易流程闭环']),
    ('语义导航（真实截图）', IMG['shot_navigation'], ['语义地点识别', '多偏好路线规划', '结构化导航卡片']),
]
for i, (name, img, lines) in enumerate(scenario):
    x = 0.75 + i * 4.2
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.4), Inches(3.85), Inches(5.8))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = RGBColor(203, 213, 225)
    if img.exists():
        slide.shapes.add_picture(str(img), Inches(x + 0.12), Inches(1.52), Inches(3.6), Inches(2.45))
    tf = slide.shapes.add_textbox(Inches(x + 0.18), Inches(4.1), Inches(3.4), Inches(2.8)).text_frame
    p = tf.paragraphs[0]
    p.text = name
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = DEEP
    for ln in lines:
        p = tf.add_paragraph()
        p.text = '• ' + ln
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT

# Slide 7: UX highlights
slide = prs.slides.add_slide(blank)
add_title(slide, '6. 交互与可解释性：让评审看见“智能处理过程”', size=33)

bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(1.25), Inches(12.0), Inches(5.95))
bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(248, 250, 252)
bg.line.color.rgb = RGBColor(226, 232, 240)

# Real screenshots: agent + dashboard + message
if IMG['shot_agent'].exists():
    slide.shapes.add_picture(str(IMG['shot_agent']), Inches(1.0), Inches(1.65), Inches(7.55), Inches(4.25))
if IMG['shot_dashboard'].exists():
    slide.shapes.add_picture(str(IMG['shot_dashboard']), Inches(8.8), Inches(1.65), Inches(3.45), Inches(2.0))
if IMG['shot_messages'].exists():
    slide.shapes.add_picture(str(IMG['shot_messages']), Inches(8.8), Inches(3.95), Inches(3.45), Inches(1.95))

caption = slide.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.2), Inches(0.95)).text_frame
caption.text = '真实页面截图：左侧为智能助手对话与结果卡片，右侧展示首页与消息中心联动。'
caption.paragraphs[0].font.size = Pt(15)
caption.paragraphs[0].font.bold = True
caption.paragraphs[0].font.color.rgb = RGBColor(30, 64, 175)

# Slide 8: Metrics & readiness
slide = prs.slides.add_slide(blank)
add_title(slide, '7. 阶段成果：可量化、可验证', size=34)

if IMG['metrics'].exists():
    slide.shapes.add_picture(str(IMG['metrics']), Inches(0.7), Inches(1.55), Inches(5.95), Inches(3.25))
if IMG['status'].exists():
    slide.shapes.add_picture(str(IMG['status']), Inches(0.7), Inches(4.95), Inches(5.95), Inches(2.2))

panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.55), Inches(5.7), Inches(5.6))
panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(255,255,255)
panel.line.color.rgb = RGBColor(203,213,225)

tf = panel.text_frame
tf.clear()
rows = [
    '代码规模快照',
    '• API 映射数量：99',
    '• Agent 相关类：44',
    '• 控制器数量：12',
    '• 前端页面视图：22',
    '',
    '构建验证',
    '• 后端编译通过：./mvnw -q -DskipTests compile',
    '• 前端构建通过：npm run build',
    '',
    '当前可优化项',
    '• 前端主包体积较大（>500kB），可通过分包进一步优化',
    '• 院审前建议补一次端到端回归脚本',
]
for i, txt in enumerate(rows):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = txt
    p.font.size = Pt(24 if i in [0,6,10] else 15)
    p.font.bold = i in [0,6,10]
    p.font.color.rgb = DEEP if i in [0,6,10] else TEXT

# Slide 9: More screenshots
slide = prs.slides.add_slide(blank)
add_title(slide, '8. 系统功能截图总览（真实页面）', size=33)
add_subtitle(slide, '覆盖失物招领、宿舍信息、个人中心等页面，保证院审展示“所见即所得”', size=15)

gallery = [
    ('失物招领', IMG['shot_lostfound']),
    ('宿舍信息', IMG['shot_dorm_info']),
    ('个人中心', IMG['shot_profile']),
]
for i, (label, shot) in enumerate(gallery):
    x = 0.75 + i * 4.2
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(3.85), Inches(5.4))
    frame.fill.solid(); frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.color.rgb = RGBColor(203, 213, 225)
    if shot.exists():
        slide.shapes.add_picture(str(shot), Inches(x + 0.12), Inches(1.78), Inches(3.6), Inches(4.5))
    tb = slide.shapes.add_textbox(Inches(x + 0.12), Inches(6.35), Inches(3.6), Inches(0.4)).text_frame
    tb.text = label
    tb.paragraphs[0].font.size = Pt(15)
    tb.paragraphs[0].font.bold = True
    tb.paragraphs[0].font.color.rgb = RGBColor(30, 64, 175)
    tb.paragraphs[0].alignment = PP_ALIGN.CENTER

# Slide 10: Demo plan
slide = prs.slides.add_slide(blank)
add_title(slide, '9. 院审现场演示建议（5-6 分钟）', size=33)
add_subtitle(slide, '按“问题提出 -> Agent处理 -> 结果落地”的顺序，让评委快速看到价值', size=15)

steps = [
    ('00:00-00:40', '开场定位', '说明目标：统一校园事务入口，降低办理门槛'),
    ('00:40-02:00', '场景一：宿舍报修', '输入自然语言报修请求，展示追问与工单结果卡片'),
    ('02:00-03:20', '场景二：二手交易', '展示筛选/收藏/消息联动，强调业务闭环'),
    ('03:20-04:40', '场景三：语义导航', '展示语义目的地识别与路线偏好（最快/安全）'),
    ('04:40-05:30', '总结与扩展', '回到架构图，说明可扩展到更多校园业务'),
]
for i,(t,stage,detail) in enumerate(steps):
    y = 1.45 + i*1.15
    left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(y), Inches(2.2), Inches(0.9))
    left.fill.solid(); left.fill.fore_color.rgb = RGBColor(219,234,254)
    left.line.color.rgb = RGBColor(147,197,253)
    left.text_frame.text = t
    left.text_frame.paragraphs[0].font.size = Pt(13)
    left.text_frame.paragraphs[0].font.bold = True
    left.text_frame.paragraphs[0].font.color.rgb = RGBColor(30,64,175)
    left.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    mid = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.25), Inches(y), Inches(2.6), Inches(0.9))
    mid.fill.solid(); mid.fill.fore_color.rgb = RGBColor(236,253,245)
    mid.line.color.rgb = RGBColor(167,243,208)
    mid.text_frame.text = stage
    mid.text_frame.paragraphs[0].font.size = Pt(14)
    mid.text_frame.paragraphs[0].font.bold = True
    mid.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.05), Inches(y), Inches(6.4), Inches(0.9))
    right.fill.solid(); right.fill.fore_color.rgb = RGBColor(248,250,252)
    right.line.color.rgb = RGBColor(203,213,225)
    right.text_frame.text = detail
    right.text_frame.paragraphs[0].font.size = Pt(13)
    right.text_frame.paragraphs[0].font.color.rgb = TEXT

# Slide 11: Risks & plan
slide = prs.slides.add_slide(blank)
add_title(slide, '10. 风险识别与答辩前优化计划', size=33)

risk_data = [
    ('已识别风险', '影响', '应对策略（院审前）'),
    ('前端打包体积偏大', '首次加载速度受影响', '按模块拆包，优先拆导航地图与图表依赖'),
    ('多模块联动路径长', '演示时容易触发边界问题', '固定院审演示账号与脚本，提前回归 3 条主链路'),
    ('部分数据依赖外部服务', '网络波动影响稳定性', '准备本地兜底数据与降级提示'),
]
x = [0.75, 4.3, 7.6]
w = [3.4, 3.1, 5.0]
for r, row in enumerate(risk_data):
    y = 1.65 + r*1.3
    for c, txt in enumerate(row):
        cell = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x[c]), Inches(y), Inches(w[c]), Inches(1.1))
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = RGBColor(30, 64, 175)
            fc = WHITE
        else:
            cell.fill.fore_color.rgb = RGBColor(248, 250, 252)
            fc = DEEP if c == 0 else TEXT
        cell.line.color.rgb = RGBColor(203, 213, 225)
        tf = cell.text_frame
        tf.text = txt
        tf.paragraphs[0].font.size = Pt(14 if r else 15)
        tf.paragraphs[0].font.bold = (r == 0 or c == 0)
        tf.paragraphs[0].font.color.rgb = fc
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

commit = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(6.0), Inches(11.9), Inches(0.9))
commit.fill.solid(); commit.fill.fore_color.rgb = RGBColor(220, 252, 231)
commit.line.color.rgb = RGBColor(134, 239, 172)
commit.text_frame.text = '答辩承诺：优先保证“演示稳定性 + 关键链路闭环”，再逐步优化性能与体验细节。'
commit.text_frame.paragraphs[0].font.size = Pt(16)
commit.text_frame.paragraphs[0].font.bold = True
commit.text_frame.paragraphs[0].font.color.rgb = RGBColor(20, 83, 45)
commit.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# Slide 12: Ending
slide = prs.slides.add_slide(blank)
add_bg(slide, IMG['library'])
add_dark_overlay(slide, 0.45)
add_title(slide, '11. 总结', left=0.9, top=1.3, size=42, color=WHITE)
add_subtitle(slide, '我们已经完成从“多入口系统”到“Agent 驱动统一入口”的关键跃迁', left=0.92, top=2.15, width=11.0, size=20, color=RGBColor(226,232,240))

sum_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(3.0), Inches(11.7), Inches(2.6))
sum_box.fill.solid(); sum_box.fill.fore_color.rgb = RGBColor(15,23,42); sum_box.fill.transparency = 0.2
sum_box.line.color.rgb = RGBColor(147,197,253)

tf = sum_box.text_frame
tf.clear()
for i, txt in enumerate([
    '关键结论',
    '• 项目具备可运行、可演示、可扩展的工程基础',
    '• Agent 机制显著提升了校园事务办理体验',
    '• 通过结构化结果与多轮追问，兼顾准确性与可解释性',
    '• 已具备院审展示竞争力，后续将持续优化稳定性与性能'
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = txt
    p.font.size = Pt(25 if i == 0 else 18)
    p.font.bold = (i == 0)
    p.font.color.rgb = WHITE

thanks = slide.shapes.add_textbox(Inches(0.95), Inches(6.15), Inches(7.0), Inches(0.7)).text_frame
thanks.text = '谢谢各位老师，欢迎批评指正'
thanks.paragraphs[0].font.size = Pt(24)
thanks.paragraphs[0].font.bold = True
thanks.paragraphs[0].font.color.rgb = RGBColor(191,219,254)

qa = slide.shapes.add_textbox(Inches(10.25), Inches(6.15), Inches(2.4), Inches(0.7)).text_frame
qa.text = 'Q & A'
qa.paragraphs[0].font.size = Pt(30)
qa.paragraphs[0].font.bold = True
qa.paragraphs[0].font.color.rgb = WHITE
qa.paragraphs[0].alignment = PP_ALIGN.RIGHT

prs.save(str(OUT))
print(OUT)
