from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

ROOT = Path('/Users/yy/code/OSP')
SRC = ROOT / '项目汇报_已完成工作与后期计划_可直接汇报版.pptx'
BG = ROOT / 'campus-frontend/public/images/campus-bg1.jpg'

TITLE = '一站式校园事务智能协同平台'
SUB = '阶段成果汇报：已完成工作与下一步计划'


def clear_slide(slide):
    spTree = slide.shapes._spTree
    for shape in list(slide.shapes):
        spTree.remove(shape._element)


def add_textbox(slide, x, y, w, h, text, size=20, bold=False, color=RGBColor(0, 0, 0), align=PP_ALIGN.LEFT):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tf


def cover_v1(slide):
    # Pure minimal white
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(255, 255, 255); bg.line.fill.background()

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(3.2), Inches(2.8), Inches(0.05))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(30, 64, 175); line.line.fill.background()

    add_textbox(slide, 1.0, 1.9, 10.8, 1.0, TITLE, size=44, bold=True, color=RGBColor(15, 23, 42))
    add_textbox(slide, 1.0, 2.75, 10.8, 0.7, SUB, size=22, bold=False, color=RGBColor(30, 64, 175))

    tf = add_textbox(slide, 1.0, 4.1, 6.0, 1.5, '姓名：\n学号：\n日期：2026 年 3 月', size=18, color=RGBColor(51, 65, 85))


def cover_v2(slide):
    # Clean business, light top band
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(248, 251, 255); bg.line.fill.background()

    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(2.0))
    top.fill.solid(); top.fill.fore_color.rgb = RGBColor(23, 79, 156); top.line.fill.background()

    add_textbox(slide, 0.95, 0.72, 11.3, 0.9, TITLE, size=38, bold=True, color=RGBColor(255, 255, 255))
    add_textbox(slide, 0.95, 1.45, 11.3, 0.5, SUB, size=18, color=RGBColor(219, 234, 254))

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(2.45), Inches(11.4), Inches(3.9))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(255, 255, 255); card.line.color.rgb = RGBColor(203, 213, 225)

    add_textbox(slide, 1.35, 3.2, 4.0, 1.5, '姓名：\n学号：\n日期：2026 年 3 月', size=20, color=RGBColor(51, 65, 85))
    add_textbox(slide, 6.2, 3.35, 5.4, 2.0, '关键词\n统一入口\n智能协同\n业务闭环', size=18, color=RGBColor(30, 64, 175), align=PP_ALIGN.CENTER)


def cover_v3(slide):
    # Photo visible but clean text on left glass card
    if BG.exists():
        slide.shapes.add_picture(str(BG), Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
    over = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    over.fill.solid(); over.fill.fore_color.rgb = RGBColor(255, 255, 255); over.fill.transparency = 0.32; over.line.fill.background()

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.85), Inches(6.2), Inches(5.9))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(19, 60, 128); panel.fill.transparency = 0.08; panel.line.color.rgb = RGBColor(147, 197, 253)

    add_textbox(slide, 1.2, 1.55, 5.6, 1.2, TITLE, size=40, bold=True, color=RGBColor(255, 255, 255))
    add_textbox(slide, 1.2, 2.85, 5.6, 0.7, SUB, size=20, color=RGBColor(219, 234, 254))
    add_textbox(slide, 1.2, 4.2, 4.8, 1.6, '姓名：\n学号：\n日期：2026 年 3 月', size=18, color=RGBColor(255, 255, 255))


def build(name, fn):
    p = Presentation(str(SRC))
    s1 = p.slides[0]
    clear_slide(s1)
    fn(s1)
    out = ROOT / name
    p.save(str(out))
    print(out)


build('项目汇报_封面方案1_极简白底.pptx', cover_v1)
build('项目汇报_封面方案2_商务蓝带.pptx', cover_v2)
build('项目汇报_封面方案3_校园背景.pptx', cover_v3)
